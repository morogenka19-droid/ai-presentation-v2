from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches
import os
import uuid
import requests
from dotenv import load_dotenv
import json

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class TopicRequest(BaseModel):
    topic: str
    template: str = "modern"

DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY")
client = OpenAI(
    api_key=DEEPSEEK_API_KEY,
    base_url="https://api.deepseek.com"
)

def get_image_url(query):
    try:
        url = f"https://source.unsplash.com/featured/?{query.replace(' ', ',')}"
        return url
    except:
        return None

"/generate"
async def generate_presentation(request: TopicRequest):
    try:
        prompt = f"""
        Составь структуру презентации на тему "{request.topic}".
        Создай 7 слайдов. Каждый слайд должен содержать заголовок и 3-4 пункта.
        Верни ответ строго в формате JSON:
        [
            {{"title": "Заголовок 1", "points": ["пункт 1", "пункт 2", "пункт 3"], "image": "ключевые слова"}},
            {{"title": "Заголовок 2", "points": ["пункт 1", "пункт 2", "пункт 3"], "image": "ключевые слова"}}
        ]
        """
        
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[
                {"role": "system", "content": "Ты помогаешь создавать структуру презентаций."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.7
        )
        
        slides_data = json.loads(response.choices[0].message.content)
        
        prs = Presentation()
        
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = request.topic
        slide.placeholders[1].text = f"Создано с помощью AI • {request.template}"
        
        for slide_info in slides_data:
            bullet_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = slide_info["title"]
            content = slide.placeholders[1]
            text_frame = content.text_frame
            text_frame.text = ""
            for point in slide_info["points"]:
                p = text_frame.add_paragraph()
                p.text = point
                p.level = 0
            
            if "image" in slide_info and slide_info["image"]:
                try:
                    img_url = get_image_url(slide_info["image"])
                    if img_url:
                        img_response = requests.get(img_url, timeout=5)
                        if img_response.status_code == 200:
                            with open("temp_img.jpg", "wb") as f:
                                f.write(img_response.content)
                            left = Inches(8)
                            top = Inches(1.5)
                            slide.shapes.add_picture("temp_img.jpg", left, top, width=Inches(3), height=Inches(3))
                            os.remove("temp_img.jpg")
                except:
                    pass
        
        filename = f"presentation_{uuid.uuid4()}.pptx"
        prs.save(filename)
        
        return FileResponse(
            path=filename,
            filename=f"Презентация_{request.topic[:20]}.pptx",
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/")
def root():
    return {"message": "AI Presentation Generator API работает!"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=10000)